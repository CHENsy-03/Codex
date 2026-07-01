import sys
sys.path.insert(0, 'C:\\Users\\35594\\.cache\\codex-runtimes\\codex-primary-runtime\\dependencies\\python')
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
from lxml import etree
import os

source_dir = 'D:\\企业微信\\文件储存\\WXWork\\1688854658857559\\Cache\\File\\2026-07'
out_dir = 'E:\\AI_Projects\\Codex\\output_images'

counties = ['威信县', '永善县', '昭阳区']

for county in counties:
    print(f'\n=== Processing {county} ===')
    src_path = os.path.join(source_dir, f'昭通市{county}_低空经济_架构思维导图版_更新版.docx')
    doc = Document(src_path)
    
    # Find Gantt chart locations
    gantt_count = 0
    
    for i, table in enumerate(doc.tables):
        is_gantt = False
        for cell in table.rows[0].cells:
            text = cell.text.strip()
            if '第' in text and ('天' in text or '周' in text):
                is_gantt = True
                break
        
        if not is_gantt or len(table.rows[0].cells) < 3:
            continue
        
        gantt_count += 1
        gantt_img_path = os.path.join(out_dir, f'{county}_gantt_{gantt_count}.png')
        
        if os.path.exists(gantt_img_path):
            # Get the table element's parent
            tbl = table._tbl
            parent = tbl.getparent()
            
            # Find previous sibling paragraph (the heading before table)
            prev = tbl.getprevious()
            while prev is not None and prev.tag != '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p':
                prev = prev.getprevious()
            
            # Add the image paragraph after the heading paragraph
            if prev is not None:
                p_img = doc.add_paragraph()
                run = p_img.add_run()
                run.add_picture(gantt_img_path, width=Cm(14))
                
                # Move this new paragraph right after the heading
                p_element = p_img._element
                prev.addnext(p_element)
                
                print(f'  Inserted Gantt {gantt_count} for {county}')
    
    # Save
    doc.save(src_path)
    print(f'  Saved {county}')

print('\nDone inserting Gantt charts!')
