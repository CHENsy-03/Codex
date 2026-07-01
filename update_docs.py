import re, os, sys
sys.path.insert(0, 'C:\\Users\\35594\\.cache\\codex-runtimes\\codex-primary-runtime\\dependencies\\python')
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import numpy as np

font_path = 'C:\\Windows\\Fonts\\msyh.ttc'
font = FontProperties(fname=font_path, size=8)
font_small = FontProperties(fname=font_path, size=6)
font_title = FontProperties(fname=font_path, size=11)

out_dir = 'E:\\AI_Projects\\Codex\\output_images'
os.makedirs(out_dir, exist_ok=True)

def set_cell_shading(cell, color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def format_table(table, total_width_cm=16):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    tblW = parse_xml(f'<w:tblW {nsdecls("w")} w:w="{int(total_width_cm*567)}" w:type="dxa"/>')
    tblPr.append(tblW)
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            vAlign = parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>')
            tcPr.append(vAlign)
            for paragraph in cell.paragraphs:
                paragraph.alignment = 1
                for run in paragraph.runs:
                    run.font.size = Pt(8)
                    run.font.name = 'Microsoft YaHei'
                    rPr = run._element.get_or_add_rPr()
                    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="Microsoft YaHei"/>')
                    rPr.append(rFonts)
            if i == 0:
                set_cell_shading(cell, '2F5496')
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = None
                        run.font.bold = True
            elif i % 2 == 0:
                set_cell_shading(cell, 'D6E4F0')
            else:
                set_cell_shading(cell, 'FFFFFF')

def generate_gantt_chart(tasks, title, filename, max_time=None):
    fig_height = max(3, len(tasks) * 0.5 + 1.5)
    fig, ax = plt.subplots(figsize=(10, fig_height))
    fig.patch.set_facecolor('white')
    ax.set_title(title, fontproperties=font_title, pad=12, fontweight='bold')
    ax.set_facecolor('#FAFBFC')
    colors = ['#1976D2', '#388E3C', '#F57C00', '#D32F2F', '#7B1FA2',
              '#0097A7', '#E64A19', '#455A64', '#5D4037', '#303F9F',
              '#00796B', '#AFB42B', '#689F38', '#FFA000', '#C2185B',
              '#512DA8', '#0288D1', '#C0CA33', '#00ACC1', '#F4511E']
    y_positions = range(len(tasks))
    for i, (task_name, periods) in enumerate(tasks):
        for period_name, (start, end) in periods.items():
            bar_color = colors[i % len(colors)]
            ax.barh(i, end - start, left=start, height=0.6,
                    color=bar_color, edgecolor='white', linewidth=0.5, alpha=0.85)
    ax.set_yticks(list(y_positions))
    ax.set_yticklabels([t[0] for t in tasks], fontproperties=font)
    ax.invert_yaxis()
    if max_time:
        ax.set_xlim(0, max_time)
    ax.grid(axis='x', alpha=0.3, linestyle='--', color='#E0E0E0')
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)
    ax.tick_params(axis='x', labelsize=7)
    ax.tick_params(axis='y', labelsize=7)
    plt.tight_layout()
    filepath = os.path.join(out_dir, filename)
    plt.savefig(filepath, dpi=150, bbox_inches='tight', facecolor='white')
    plt.close()
    print(f'  Saved: {filename}')
    return filepath

# Process all 3 documents
counties_list = ['威信县', '永善县', '昭阳区']
source_dir = 'D:\\企业微信\\文件储存\\WXWork\\1688854658857559\\Cache\\File\\2026-07'

for county in counties_list:
    print(f'\n=== Processing {county} ===')
    src_path = os.path.join(source_dir, f'昭通市{county}_低空经济_架构思维导图版.docx')
    doc = Document(src_path)
    
    # Format all tables
    table_count = 0
    for i, table in enumerate(doc.tables):
        try:
            format_table(table, 16)
            table_count += 1
        except Exception as e:
            print(f'  Warning: Table {i}: {e}')
    print(f'  Formatted {table_count} tables')
    
    # Save updated document
    output_path = os.path.join(source_dir, f'昭通市{county}_低空经济_架构思维导图版_更新版.docx')
    doc.save(output_path)
    print(f'  Saved: {output_path}')

print('\n=== Done formatting documents! ===')
