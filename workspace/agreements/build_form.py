import sys,os;sys.stdout.reconfigure(encoding="utf-8")
from docx import Document;from docx.shared import Pt,Cm,Inches;from docx.enum.text import WD_ALIGN_PARAGRAPH as A
doc=Document();S=doc.styles["Normal"];S.font.name="SimSun";S.font.size=Pt(11)
def CT(cell,text,bold=False,size=11,align="left"):
 p=cell.paragraphs[0];p.alignment={"left":0,"center":1}[align]
 r=p.add_run(text);r.font.size=Pt(size);r.font.name="SimSun"
 if bold:r.bold=True
 p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
def MT(c,h1,h2,v1,v2):
 r=c.add_row();CT(r.cells[0],h1,True,11,"center");CT(r.cells[1],v1);CT(r.cells[2],h2,True,11,"center");CT(r.cells[3],v2)
def RT(c,h,v,col=3):
 r=c.add_row();CT(r.cells[0],h,True,11,"center");r.cells[0].merge(r.cells[3]);CT(r.cells[0],v)
def SH(c,text):
 r=c.add_row();CT(r.cells[0],text,True,12,"center")
 r.cells[0].merge(r.cells[3])

# Title
p=doc.add_paragraph();p.alignment=A.CENTER;r=p.add_run("宾航私董会申请表");r.bold=True;r.font.size=Pt(22);r.font.name="SimHei"
p.paragraph_format.space_after=Pt(4)
p2=doc.add_paragraph();p2.alignment=A.CENTER;r2=p2.add_run("（请如实填写，* 为必填项）");r2.font.size=Pt(10);r2.font.color.rgb=None
p2.paragraph_format.space_after=Pt(12)

# Table 1: Personal Info
t1=doc.add_table(rows=0,cols=4);t1.style="Table Grid"
t1.autofit=True
SH(t1,"个人基本信息")
MT(t1,"姓名","性别","","")
MT(t1,"民族","文化程度","","")
MT(t1,"籍贯","党派","","")
MT(t1,"出生年月","邮箱","","")
MT(t1,"联系电话","社会职务","","")
r=t1.add_row();CT(r.cells[0],"个人简介",True,11,"center");r.cells[0].merge(r.cells[3]);CT(r.cells[0],"")
r=t1.add_row();CT(r.cells[0],"私董会职务意向",True,11,"center");r.cells[0].merge(r.cells[3])
CT(r.cells[0],"\u25a1 普通会员  \u25a1 理事会员  \u25a1 常务理事  \u25a1 副会长  \u25a1 常务副会长  \u25a1 会长")

doc.add_paragraph()

# Table 2: Enterprise Info
t2=doc.add_table(rows=0,cols=4);t2.style="Table Grid"
t2.autofit=True
SH(t2,"企业基本情况")
RT(t2,"名称（中）","")
RT(t2,"名称（英）","")
MT(t2,"工商登记号","成立日期","","")
MT(t2,"经济性质","所属行业","","")
MT(t2,"注册资金","资产总额","","")
MT(t2,"公司净资产","上年销售额","","")
MT(t2,"上年纳税额","职工人数","","")
MT(t2,"分支机构","所属行业(分支)","","")
RT(t2,"通讯地址","")

# Footer
doc.add_paragraph()
p3=doc.add_paragraph();r3=p3.add_run("申请人签名：__________________    日期：______年______月______日")
r3.font.size=Pt(10)
p4=doc.add_paragraph();p4.alignment=A.CENTER;r4=p4.add_run("宾航私董会秘书处 制表 \u00b7 本表复印有效")
r4.font.size=Pt(9);r4.font.color.rgb=None

# Save
out=r"E:\AI_Projects\Codex\协议修改与制作\宾航私董会申请表.docx"
doc.save(out)
print("Saved: "+out)
print("DONE!")