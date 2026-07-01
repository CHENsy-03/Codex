import sys, os, re
sys.path.insert(0, r"C:\Users\35594\.cache\codex-runtimes\codex-primary-runtime\dependencies\python")
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

font_path = r"C:\Windows\Fonts\msyh.ttc"
out_dir = r"E:\AI_Projects\Codex\output_images\v3"
os.makedirs(out_dir, exist_ok=True)
source_dir = r"D:\企业微信\文件储存\WXWork\1688854658857559\Cache\File\2026-07"

# Heartbeat counter
hb = 0
def heartbeat():
    global hb
    hb += 1
    if hb % 5 == 0:
        print(f"[heartbeat: processed {hb} items]")

# === MIND MAP DATA ===
counties_mindmaps = {
    "威信县": [
        {"title":"威信县低空经济三层架构","nodes":{
            "第一层\n基建底座":["空域规划与划设","数字底座绘制","感知网络铺建","5G通信建设","气象设施部署"],
            "第二层\n枢纽运营":["五大功能中心","指挥调度中心","运维保障中心","培训展示中心"],
            "第三层\n服务网络":["红色旅游观光","特色农业服务","矿区物资运输","城市治理应急","飞行器维保"]
        },"fname":"威信县_mindmap_1_layer_v3.png"},
        {"title":"威信县勘测组织架构","nodes":{
            "总指挥部":["项目经理1名","统筹调度协调"],
            "10个勘测片区":["每区8人勘测队","覆盖全县乡镇","动态调度机制"],
            "技术支持组":["GIS工程师","数据分析师","设备维护"],
            "后勤保障组":["车辆调度","物资供应","住宿餐饮"],
            "质量控制组":["数据质检","进度监控","安全监督"]
        },"fname":"威信县_mindmap_2_org_v3.png"},
        {"title":"威信县低空运行管理云平台","nodes":{
            "数据底座层":["三维地图数据","设备点位数据","气象数据融合","统一数据接口"],
            "业务应用层":["低空智能网联","航线审批管理","实时飞行监控","数字孪生系统"],
            "空管保障层":["空域管理","流量监控","冲突解算","安全审计"],
            "配套服务层":["用户管理","计费系统","数据分析","报告生成"]
        },"fname":"威信县_mindmap_3_platform_v3.png"}
    ],
    "永善县": [
        {"title":"永善县低空经济三层架构","nodes":{
            "第一层\n基建底座":["空域规划与划设","三维数字底座","感知网络铺建","通信网络建设"],
            "第二层\n枢纽运营":["运营服务中心","物流集散中心","维修保障中心","培训展示中心"],
            "第三层\n服务网络":["农业植保服务","物流运输","旅游观光","城市治理"]
        },"fname":"永善县_mindmap_1_layer_v3.png"},
        {"title":"永善县勘测组织架构","nodes":{
            "总指挥部":["项目经理","技术总监","安全监督"],
            "16个勘测片区":["每区6-7人","动态增援机制","三阶段突击策略"],
            "后勤技术组":["设备保障","数据管理","技术支撑"],
            "质量控制组":["质检专员","进度跟踪","安全巡查"]
        },"fname":"永善县_mindmap_2_org_v3.png"},
        {"title":"永善县低空运行管理云平台","nodes":{
            "数据底座层":["三维地图导入","设备数据接入","气象数据融合","标准化处理"],
            "核心业务层":["身份识别","航线审批","冲突解算","实时监控","数字孪生"],
            "空管保障层":["气象接入","雷达接入","5G-A感知","告警联动"],
            "业务支撑层":["物流调度","植保管理","旅游服务","应急指挥"]
        },"fname":"永善县_mindmap_3_platform_v3.png"}
    ],
    "昭阳区": [
        {"title":"昭阳区低空经济三层架构","nodes":{
            "第一层\n基建底座":["空域规划与划设","数字底座绘制","5G-A通信网络","雷达感知网络"],
            "第二层\n枢纽运营":["运营服务中心","物流集散中心","维保中心","数据中心"],
            "第三层\n服务网络":["即时物流","城市治理","载人观光","飞行器维保"]
        },"fname":"昭阳区_mindmap_1_layer_v3.png"},
        {"title":"昭阳区勘测组织架构","nodes":{
            "指挥中心":["项目总监","技术主管","安全监督"],
            "20个勘测片区":["每区10人","覆盖全区乡镇","动态调度"],
            "技术支撑组":["GIS分析","数据处理","设备维护"],
            "后勤保障组":["车辆调度","物资供应","应急保障"],
            "质量监督组":["数据验收","进度管控","安全巡查"]
        },"fname":"昭阳区_mindmap_2_org_v3.png"},
        {"title":"昭阳区低空运行管理云平台","nodes":{
            "数据底座层":["三维地图数据","设备数据接入","气象数据融合","标准化接口"],
            "核心业务层":["身份识别","航线审批","冲突解算","实时监控","数字孪生"],
            "空管保障层":["态势感知","告警联动","应急响应","安全审计"],
            "场景应用层":["物流调度","城市治理","载人观光","维保管理"]
        },"fname":"昭阳区_mindmap_3_platform_v3.png"}
    ]
}

print("=== Generating Projection Mind Maps ===")
for county, maps in counties_mindmaps.items():
    print(f"  {county}: {len(maps)} maps")
    for mm in maps:
        fp = os.path.join(out_dir, mm["fname"])
        fig, ax = plt.subplots(figsize=(22, 12))
        fig.patch.set_facecolor("#FAFBFC")
        ax.set_xlim(0, 22)
        ax.set_ylim(0, 12)
        ax.axis("off")
        
        ftitle = FontProperties(fname=font_path, size=36, weight="bold")
        fnode = FontProperties(fname=font_path, size=26, weight="bold")
        fsub = FontProperties(fname=font_path, size=20)
        
        colors = ["#1565C0","#2E7D32","#E65100","#C62828","#6A1B9A","#00838F","#558B2F"]
        box_colors = ["#E3F2FD","#E8F5E9","#FFF3E0","#FFEBEE","#F3E5F5","#E0F7FA","#F1F8E9"]
        
        ax.text(11, 11.6, mm["title"], ha="center", va="center", fontproperties=ftitle, color="#1a1a2e")
        ax.plot([3, 19], [11.0, 11.0], "-", color="#1a1a2e", lw=2.5)
        
        nodes = mm["nodes"]
        num = len(nodes)
        if num == 0:
            plt.savefig(fp, dpi=300, bbox_inches="tight")
            plt.close()
            continue
        spacing = 22.0 / (num + 1)
        for i, (bn, subs) in enumerate(nodes.items()):
            x = spacing * (i + 1)
            yt = 10.4
            c = colors[i % len(colors)]
            bc = box_colors[i % len(box_colors)]
            
            ax.text(x, yt, bn, ha="center", va="center", fontproperties=fnode, color="white",
                    bbox=dict(boxstyle="round,pad=0.7", facecolor=c, edgecolor=c, lw=2))
            for j, sub in enumerate(subs):
                ys = yt - (j + 1) * 0.9 - 0.5
                if ys < 0.4: break
                ax.plot([x, x], [ys+0.4, ys-0.4], "-", color=c, lw=1.5, alpha=0.3)
                ax.text(x, ys, sub, ha="center", va="center", fontproperties=fsub, color="#222222",
                        bbox=dict(boxstyle="round,pad=0.4", facecolor=bc, edgecolor=c, lw=1.5))
        plt.tight_layout()
        plt.subplots_adjust(top=0.93, bottom=0.05)
        plt.savefig(fp, dpi=300, bbox_inches="tight", facecolor="white")
        plt.close()
        print(f"    Saved: {mm['fname']}")
        heartbeat()

# === GANTT CHARTS ===
def parse_period(ps):
    ps = ps.strip()
    m = re.search(r"第(\d+)[-~](\d+)(天|周)", ps)
    if m:
        s, e = int(m.group(1)), int(m.group(2))
        if m.group(3) == "周":
            s, e = s*7, e*7
        return (s, e)
    m = re.search(r"第(\d+)(天|周)", ps)
    if m:
        v = int(m.group(1))
        if m.group(1) == "周": v*=7
        return (max(1,v-3), v+3)
    if "长期" in ps or "贯穿" in ps: return None
    return None

print("\n=== Generating Gantt Charts ===")
counties_gantt = {}
for county, fname in [("威信县","昭通市威信县_低空经济_架构思维导图版.docx"),
                        ("永善县","昭通市永善县_低空经济_架构思维导图版.docx"),
                        ("昭阳区","昭通市昭阳区_低空经济_架构思维导图版.docx")]:
    fp = os.path.join(source_dir, fname)
    doc = Document(fp)
    print(f"  {county}: {len(doc.tables)} tables")
    gc = 0
    gantt_list = []
    for i, table in enumerate(doc.tables):
        if len(table.rows) < 2 or len(table.columns) < 3: continue
        tlabels = []
        for cell in table.rows[0].cells:
            t = cell.text.strip()
            if "第" in t and ("天" in t or "周" in t): tlabels.append(t)
        if len(tlabels) < 2: continue
        
        tasks = []
        for j, row in enumerate(table.rows):
            if j == 0: continue
            cells = [c.text.strip() for c in row.cells]
            tn = cells[0] if cells else ""
            if not tn: continue
            periods = []
            for k, tl in enumerate(tlabels):
                if k+1 < len(cells):
                    ct = cells[k+1]
                    if len([c for c in ct if c in "█▉▊▋▌▍▎▏"]) > 0 or (ct.strip() and "—" not in ct):
                        pp = parse_period(tl)
                        if pp: periods.append(pp)
            if periods: tasks.append((tn, periods))
        
        if tasks:
            gc += 1
            title = f"{county} - 第{gc}阶段实施甘特图"
            fn = f"{county}_gantt_{gc}_v3.png"
            
            fh = max(3.5, len(tasks) * 0.55 + 1.5)
            fig, ax = plt.subplots(figsize=(14, fh))
            fig.patch.set_facecolor("white")
            ft = FontProperties(fname=font_path, size=14, weight="bold")
            ax.set_title(title, fontproperties=ft, pad=12, color="#1a1a2e", fontsize=16)
            ax.set_facecolor("#FAFBFC")
            
            colors = ["#1976D2","#388E3C","#F57C00","#D32F2F","#7B1FA2","#0097A7","#E64A19","#455A64","#5D4037","#303F9F"]
            ypos = range(len(tasks))
            mx = 0
            for k, (tn_, periods) in enumerate(tasks):
                for s, e in periods:
                    if s is not None and e is not None:
                        ax.barh(k, e-s, left=s, height=0.55, color=colors[k%len(colors)], edgecolor="white", lw=0.5, alpha=0.85)
                        if e > mx: mx = e
            
            ax.set_yticks(list(ypos))
            fl = FontProperties(fname=font_path, size=10)
            ax.set_yticklabels([t[0] for t in tasks], fontproperties=fl)
            ax.invert_yaxis()
            if mx == 0: mx = 100
            ax.set_xlim(0, mx*1.05)
            step = 5 if mx <= 30 else (10 if mx <= 100 else (20 if mx <= 150 else 30))
            ticks = list(range(0, int(mx) + step, step))
            fs = FontProperties(fname=font_path, size=9)
            ax.set_xticks(ticks)
            ax.set_xticklabels([str(t)+"天" for t in ticks], fontproperties=fs)
            ax.grid(axis="x", alpha=0.3, linestyle="--", color="#E0E0E0")
            ax.spines["top"].set_visible(False)
            ax.spines["right"].set_visible(False)
            plt.tight_layout()
            fpath = os.path.join(out_dir, fn)
            plt.savefig(fpath, dpi=200, bbox_inches="tight", facecolor="white")
            plt.close()
            print(f"    Gantt: {fn}")
            gantt_list.append({"title": title, "fname": fn, "path": fpath, "table_index": i})
            heartbeat()
    counties_gantt[county] = gantt_list

# === FORMAT TABLES & SAVE ===
def set_shading(cell, color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def fmt_table(table, w=16):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = parse_xml(f'<w:tblPr {nsdecls("w")}/>')
        tbl.insert(0, tblPr)
    for ex in tblPr.findall(qn("w:tblW")):
        tblPr.remove(ex)
    tblPr.append(parse_xml(f'<w:tblW {nsdecls("w")} w:w="{int(w*567)}" w:type="dxa"/>'))
    
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            for ex in tcPr.findall(qn("w:vAlign")):
                tcPr.remove(ex)
            tcPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))
            
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(8)
                    r.font.name = "Microsoft YaHei"
                    rPr = r._element.get_or_add_rPr()
                    for ex in rPr.findall(qn("w:rFonts")):
                        rPr.remove(ex)
                    rPr.append(parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="Microsoft YaHei"/>'))
                    if ri == 0:
                        r.font.bold = True
                        for ex in rPr.findall(qn("w:color")):
                            rPr.remove(ex)
                        rPr.append(parse_xml(f'<w:color {nsdecls("w")} w:val="FFFFFF"/>'))
            if ri == 0:
                set_shading(cell, "1F4E79")
            elif ri % 2 == 0:
                set_shading(cell, "D6E4F0")
            else:
                set_shading(cell, "FFFFFF")

print("\n=== Formatting Tables & Saving ===")
for county, fname in [("威信县","昭通市威信县_低空经济_架构思维导图版.docx"),
                        ("永善县","昭通市永善县_低空经济_架构思维导图版.docx"),
                        ("昭阳区","昭通市昭阳区_低空经济_架构思维导图版.docx")]:
    fp = os.path.join(source_dir, fname)
    doc = Document(fp)
    tc = 0
    for i, table in enumerate(doc.tables):
        try:
            fmt_table(table, 16)
            tc += 1
        except Exception as e:
            print(f"    Table {i} warning: {e}")
    print(f"  {county}: formatted {tc} tables")
    
    outpath = os.path.join(source_dir, f"昭通市{county}_低空经济_架构思维导图版_更新版V3.docx")
    doc.save(outpath)
    print(f"  Saved: {outpath}")
    heartbeat()

print("\n=== DONE! ===")
