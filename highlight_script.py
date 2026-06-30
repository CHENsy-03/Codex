import os
import sys
import copy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def _get_or_create_rPr(run_element):
    rPr = run_element.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        run_element.insert(0, rPr)
    return rPr

def _ensure_highlight(run_element):
    rPr = _get_or_create_rPr(run_element)
    b = rPr.find(qn('w:b'))
    if b is None:
        b = OxmlElement('w:b')
        rPr.append(b)
    u = rPr.find(qn('w:u'))
    if u is None:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)
    else:
        u.set(qn('w:val'), 'single')
    color = rPr.find(qn('w:color'))
    if color is None:
        color = OxmlElement('w:color')
        color.set(qn('w:val'), '000000')
        rPr.append(color)
    else:
        color.set(qn('w:val'), '000000')

def _make_run_element(p_elem, text, rPr_template=None, highlight=False):
    r = OxmlElement('w:r')
    if rPr_template is not None:
        rPr = copy.deepcopy(rPr_template)
        r.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    r.append(t)
    if highlight:
        _ensure_highlight(r)
    p_elem.append(r)
    return r

def process_paragraph(para, keywords):
    p_elem = para._p
    runs = p_elem.findall(qn('w:r'))
    if not runs:
        return
    full_text = ''
    char_to_rPr = []
    for r in runs:
        t_elem = r.find(qn('w:t'))
        text = t_elem.text if t_elem is not None else ''
        rPr = r.find(qn('w:rPr'))
        for c in text:
            full_text += c
            char_to_rPr.append(rPr)
    if not full_text.strip():
        return
    matches = []
    for kw in keywords:
        start = 0
        while True:
            idx = full_text.find(kw, start)
            if idx == -1:
                break
            matches.append((idx, idx + len(kw)))
            start = idx + 1
    if not matches:
        return
    matches.sort()
    merged = []
    for m in matches:
        if merged and m[0] <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], m[1]))
        else:
            merged.append(m)
    split_points = sorted(set([0, len(full_text)] + [s for m in merged for s in m]))
    segments = []
    for i in range(len(split_points) - 1):
        s = split_points[i]
        e = split_points[i + 1]
        text = full_text[s:e]
        if not text:
            continue
        is_highlight = any(ms <= s < me for ms, me in merged)
        rPr_template = char_to_rPr[s] if s < len(char_to_rPr) else None
        segments.append((text, rPr_template, is_highlight))
    for r in runs:
        p_elem.remove(r)
    for text, rPr_template, is_highlight in segments:
        _make_run_element(p_elem, text, rPr_template, is_highlight)

def main():
    keywords = [
        '投资', '投资款', '投资总额', '投资合作', '生产建设专项资金',
        '进价', '价格保护', '最低价格', '基准价格', '价格标准',
        '采购价格', '价格优惠', '价格调整', '价格审查',
        '进货数量', '采购数量', '采购目标', '采购计划',
        '采购订单', '采购量',
        '一期采购', '二期采购',
        '实际销售数量',
        '股权比例', '股权调整', '初始股权',
        '股权交割', '股权转让', '持股比例', '对赌',
        '合作保护期', '排他性合作期限',
        '合作期限', '保护期', '排他性义务', '排他性',
        '有效期',
        '知识产权', '背景知识产权', '共有知识产权',
        '专利', '著作权', '技术成果',
        '违约金', '保证金',
    ]
    src = r'E:\AI_Projects\Codex\source_agreement.docx'
    dst = r'C:\Users\35594\Desktop\板卡与通信模块\战略合作框架协议.docx'
    doc = Document(src)
    # Diagnostic for para 178
    p178 = doc.paragraphs[178]
    t178 = p178.text
    proc_count = 0
    for pi, para in enumerate(doc.paragraphs):
        process_paragraph(para, keywords)
        if pi == 178:
            proc_count = 1
    print(f"DBG: Processed para 178? {'YES' if proc_count else 'NO'}")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    process_paragraph(para, keywords)
    doc.save(dst)
    print('Done')

if __name__ == '__main__':
    main()
